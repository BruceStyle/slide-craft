#!/usr/bin/env python3
"""从旧 PPT 提取素材:逐页文字 + 全部图片 + 一张带文件名的缩略索引图。

用法:
  python extract_pptx.py old.pptx                # 输出到 ./old_extract/
  python extract_pptx.py old.pptx --out 素材盘点

输出:
  text.txt     逐页文字
  media/       PPT 内嵌的全部图片(原样解出)
  index.jpg    过滤掉小图标后的缩略索引图,交给用户圈选要用的素材
"""
import argparse
import pathlib
import sys
import zipfile

from PIL import Image, ImageDraw
from pptx import Presentation

MIN_SIDE = 380      # 短于这个像素的多半是图标
MIN_BYTES = 40_000  # 小于这个体积的多半是图标
THUMB = 240
COLS = 5


def extract_text(pptx_path, out):
    prs = Presentation(str(pptx_path))
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame]
        lines.append(f"[P{i:02d}] " + " / ".join(t.replace("\n", " ") for t in texts if t))
    (out / "text.txt").write_text("\n".join(lines), encoding="utf-8")
    return len(lines)


def extract_media(pptx_path, out):
    media = out / "media"
    media.mkdir(exist_ok=True)
    files = []
    with zipfile.ZipFile(pptx_path) as z:
        for name in z.namelist():
            if name.startswith("ppt/media/") and not name.endswith("/"):
                target = media / pathlib.Path(name).name
                target.write_bytes(z.read(name))
                files.append(target)
    return sorted(files)


def build_index(files, out):
    keep = []
    for f in files:
        if f.stat().st_size < MIN_BYTES:
            continue
        try:
            with Image.open(f) as im:
                if max(im.size) < MIN_SIDE:
                    continue
                im = im.convert("RGB")
                im.thumbnail((THUMB, THUMB))
                keep.append((f.name, im.copy()))
        except Exception:
            continue  # emf/wmf 等 Pillow 打不开的格式,跳过
    if not keep:
        return 0
    cell_w, cell_h = THUMB + 16, THUMB + 34
    rows = (len(keep) + COLS - 1) // COLS
    sheet = Image.new("RGB", (COLS * cell_w, rows * cell_h), "white")
    draw = ImageDraw.Draw(sheet)
    for k, (name, im) in enumerate(keep):
        x, y = (k % COLS) * cell_w + 8, (k // COLS) * cell_h + 8
        sheet.paste(im, (x + (THUMB - im.width) // 2, y + (THUMB - im.height) // 2))
        draw.text((x, y + THUMB + 6), name, fill="black")
    sheet.save(out / "index.jpg", quality=85)
    return len(keep)


def main():
    ap = argparse.ArgumentParser(description="旧 PPT 素材提取")
    ap.add_argument("pptx", help="旧 PPT 文件(.pptx)")
    ap.add_argument("--out", help="输出目录")
    a = ap.parse_args()

    src = pathlib.Path(a.pptx).expanduser()
    if not src.is_file():
        sys.exit(f"找不到文件: {src}")
    if src.suffix.lower() != ".pptx":
        sys.exit("只支持 .pptx;旧版 .ppt 请先用 PowerPoint 或 Keynote 另存为 .pptx。")
    out = pathlib.Path(a.out) if a.out else pathlib.Path(src.stem + "_extract")
    out.mkdir(parents=True, exist_ok=True)

    pages = extract_text(src, out)
    files = extract_media(src, out)
    shown = build_index(files, out)
    print(f"{pages} 页文字 -> {out / 'text.txt'}")
    print(f"{len(files)} 张图片 -> {out / 'media'}")
    print(f"索引图含 {shown} 张(已滤掉小图标) -> {out / 'index.jpg'}" if shown else "没有够大的图片,未生成索引图")


if __name__ == "__main__":
    main()
